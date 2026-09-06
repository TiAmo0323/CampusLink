from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(__file__).resolve().parents[1]
out=root/"CampusLink项目开发进展报告_20260905.docx"
assets=root/"report_assets"; ui=assets/"progress_20260905"
BLUE="2457A6"; DARK="17365D"; LIGHT="EAF2F8"; PALE="F5F8FC"; GREEN="2E7D32"; ORANGE="D97706"; GRAY="666666"
d=Document(); s=d.sections[0]
s.page_width=Cm(21); s.page_height=Cm(29.7); s.top_margin=Cm(2.1); s.bottom_margin=Cm(1.9); s.left_margin=Cm(2.1); s.right_margin=Cm(1.9)
for st in d.styles:
    if hasattr(st,"font"):
        st.font.name="Microsoft YaHei"; st._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),"微软雅黑")
d.styles["Normal"].font.size=Pt(10.5); d.styles["Normal"].paragraph_format.line_spacing=1.35; d.styles["Normal"].paragraph_format.space_after=Pt(5)
for n,z,c in [("Heading 1",18,DARK),("Heading 2",14,BLUE)]:
    d.styles[n].font.size=Pt(z); d.styles[n].font.bold=True; d.styles[n].font.color.rgb=RGBColor.from_string(c)

def R(p,t,b=False,c=None,z=None):
    r=p.add_run(t); r.bold=b; r.font.name="Microsoft YaHei"; r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),"微软雅黑")
    if c:r.font.color.rgb=RGBColor.from_string(c)
    if z:r.font.size=Pt(z)
    return r
def shade(cell,c):
    pr=cell._tc.get_or_add_tcPr(); x=pr.find(qn("w:shd"))
    if x is None:x=OxmlElement("w:shd");pr.append(x)
    x.set(qn("w:fill"),c)
def marg(cell,v=110):
    pr=cell._tc.get_or_add_tcPr(); x=OxmlElement("w:tcMar");pr.append(x)
    for k in ("top","start","bottom","end"):
        y=OxmlElement("w:"+k);y.set(qn("w:w"),str(v));y.set(qn("w:type"),"dxa");x.append(y)
def H(t,level=1):
    p=d.add_paragraph(style="Heading "+str(level));p.paragraph_format.keep_with_next=True;p.add_run(t)
def P(t):
    p=d.add_paragraph();p.paragraph_format.first_line_indent=Cm(.74);R(p,t)
def B(t):
    p=d.add_paragraph(style="List Bullet");p.paragraph_format.left_indent=Cm(.8);p.paragraph_format.first_line_indent=Cm(-.35);R(p,t)
def T(head,rows,widths=None):
    t=d.add_table(rows=1,cols=len(head));t.style="Table Grid";t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,v in enumerate(head):t.cell(0,i).text=v
    for vals in rows:
        cells=t.add_row().cells
        for i,v in enumerate(vals):cells[i].text=str(v)
    for ri,row in enumerate(t.rows):
        for ci,cell in enumerate(row.cells):
            marg(cell);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri==0:
                shade(cell,BLUE)
                for p in cell.paragraphs:
                    for r in p.runs:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
            elif ri%2==0:shade(cell,PALE)
            if widths:cell.width=Cm(widths[ci])
    return t
def cap(t):
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(8);R(p,t,c=GRAY,z=9)
def pic(path,w,label):
    if path.exists():
        p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(path),width=Cm(w));cap(label)
def page_num(p):
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER;R(p,"第 ",c=GRAY,z=9);r=p.add_run()
    a=OxmlElement("w:fldChar");a.set(qn("w:fldCharType"),"begin");b=OxmlElement("w:instrText");b.text=" PAGE ";c=OxmlElement("w:fldChar");c.set(qn("w:fldCharType"),"end");r._r.extend([a,b,c]);R(p," 页",c=GRAY,z=9)

hp=s.header.paragraphs[0];hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT;R(hp,"CampusLink 校园互助平台｜项目开发进展报告",c=GRAY,z=9);page_num(s.footer.paragraphs[0])

p=d.add_paragraph();p.paragraph_format.space_before=Pt(65);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;R(p,"CampusLink",True,BLUE,34)
p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;R(p,"校园互助平台",True,DARK,22)
p=d.add_paragraph();p.paragraph_format.space_before=Pt(18);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;R(p,"项目开发进展报告",True,DARK,28)
p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;R(p,"第三阶段 · 第 2 周节点前进展说明",c=GRAY,z=13)
d.add_paragraph()
ct=d.add_table(rows=5,cols=2);ct.style="Table Grid";ct.alignment=WD_TABLE_ALIGNMENT.CENTER
for row,vals in zip(ct.rows,[("项目名称","CampusLink 校园互助平台"),("报告阶段","第三阶段开发收尾阶段"),("当前开发完成度","约 95%"),("计划完成日期","2026 年 9 月 13 日前"),("报告日期","2026 年 9 月 5 日")]):
    row.cells[0].text,row.cells[1].text=vals;shade(row.cells[0],LIGHT)
    for cell in row.cells:marg(cell,150);cell.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
    row.cells[0].paragraphs[0].runs[0].bold=True
p=d.add_paragraph();p.paragraph_format.space_before=Pt(55);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;R(p,"CampusLink 项目组",c=GRAY,z=11)
d.add_page_break()

H("一、项目进展概况")
P("截至 2026 年 9 月 5 日，CampusLink 校园互助平台的核心功能开发、主要业务流程、Web 前端页面和 Android 端工程适配已基本完成，当前开发完成度约为 95%。项目已进入开发收尾阶段，剩余工作主要集中在少量界面细节、边界处理、配置整理和交付材料完善。")
P("正式集中测试工作尚未启动，当前正式测试进度为 0%。开发过程中进行的编译检查、接口联调和页面自检用于支持编码实施，不计入第三阶段的正式测试工作。正式测试将在剩余开发项冻结后统一开展。")
q=T(["关键问题","当前回答"],[("1）当前开发进度","约 95%，已接近完成"),("2）能否在第三阶段第 2 周（9 月 13 日前）完成全部开发和测试","是"),("3）当前测试进展","正式测试工作尚未开展（0%）")],[8.5,7])
for r in q.cell(2,1).paragraphs[0].runs:r.bold=True;r.font.color.rgb=RGBColor.from_string(GREEN);r.font.size=Pt(14)
H("1.1 开发进度分解",2)
pt=T(["工作模块","完成度","进度示意","状态"],[
("需求理解与总体设计落实","100%","■■■■■■■■■■","已完成"),("后端核心业务开发","98%","■■■■■■■■■□","基本完成"),
("Web 前端功能与交互","96%","■■■■■■■■■□","基本完成"),("数据库结构与迁移脚本","98%","■■■■■■■■■□","基本完成"),
("Android 端工程与移动适配","90%","■■■■■■■■■□","收尾中"),("部署配置与项目文档","90%","■■■■■■■■■□","收尾中"),
("正式测试工作","0%","□□□□□□□□□□","尚未启动")],[5.2,2,5.4,2.5])
for r in pt.rows[-1].cells[1].paragraphs[0].runs+pt.rows[-1].cells[3].paragraphs[0].runs:r.font.color.rgb=RGBColor.from_string(ORANGE)
H("1.2 当前阶段判断",2)
for x in ["核心业务闭环已经形成，可完成注册登录、任务发布、申请接取、履约提交、验收结算和评价。","技能互助、积分信用、举报治理、消息通知和管理统计等扩展模块已经实现。","桌面端与移动端主要页面已完成响应式适配，Android 工程已形成可安装的调试版本。","当前重点已由大规模功能编码转为开发收尾、正式测试和交付整理。"]:B(x)
d.add_page_break()

H("二、9 月 13 日前完成情况判断")
call=d.add_table(rows=1,cols=1);call.alignment=WD_TABLE_ALIGNMENT.CENTER;shade(call.cell(0,0),LIGHT);marg(call.cell(0,0),220)
p=call.cell(0,0).paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;R(p,"是。预计可以在第三阶段第 2 周（2026 年 9 月 13 日前）完成全部开发和测试工作。",True,GREEN,14)
P("系统核心功能和主要页面已经开发完成，剩余开发量约为 5%；后续任务边界明确，可以按“开发冻结—功能测试—专项测试—缺陷回归—交付归档”推进。只要后续需求保持稳定并及时关闭高优先级缺陷，当前工期具备可执行性。")
H("2.1 后续工作计划",2)
T(["时间","工作安排","预期成果"],[
("9 月 6 日—8 日","完成剩余约 5% 的开发工作","统一界面细节、异常提示、配置和文档；冻结功能范围"),
("9 月 9 日—10 日","开展功能与接口测试","覆盖账号、任务、订单、技能、积分、评价、举报和管理端主流程"),
("9 月 11 日","开展专项测试","并发一致性、性能、安全边界、兼容性与 Android 真机检查"),
("9 月 12 日","缺陷修复与全量回归","关闭高优先级问题，复核关键数据和端到端业务闭环"),
("9 月 13 日前","完成交付","整理测试记录、部署说明、用户说明和最终版本")],[3.2,5,7.2])
H("2.2 进度保障措施",2)
for x in ["9 月 8 日前完成需求与功能冻结，新增需求统一转入后续迭代。","按主业务流程优先安排测试，发现阻断性问题立即修复并回归。","Web 端与 Android 端共用核心前端代码，统一修复交互和接口问题。","每日检查剩余事项和缺陷状态，保证高优先级问题不过夜积压。"]:B(x)
H("2.3 当前技术实现概览",2);pic(assets/"architecture.png",15.5,"图 1  CampusLink 当前系统架构")
P("系统采用前后端分离架构：后端以 Spring Boot 提供 REST 接口和业务事务，前端以 Vue 3 与 Element Plus 构建页面，Android 端通过 Capacitor 复用移动适配后的前端能力；数据库使用 MySQL，并通过 Flyway 管理结构演进。")
d.add_page_break()

H("三、已经完成的开发工作和成果")
sections=[
("3.1 账号与安全基础能力",["完成用户注册、登录、退出、令牌刷新、个人资料和账号状态管理。","完成学生用户与管理员角色区分，以及受保护接口的权限校验。","完成登录失败控制、令牌版本失效和会话管理等安全基础能力。","完成头像、任务凭证等文件上传入口及文件类型、大小限制处理。"]),
("3.2 微任务核心业务闭环",["完成任务发布、任务广场、关键字与分类筛选、任务详情和我的任务。","完成报名截止时间、任务执行时间、奖励积分和最低信用分等发布约束。","完成申请接取、发布者查看申请人、选择接取者和订单生成。","完成开始执行、成果说明与凭证提交、驳回重提、验收通过和异常处理。","完成取消、超时、重新指派和历史记录保留等边界流程。"]),
("3.3 技能互助、积分信用与评价",["完成技能分类、用户可提供技能、技能需求和熟练度维护。","完成基于技能、信用、评价和活跃度的匹配推荐及分页展示。","完成技能互助申请、双方确认和互助状态流转。","完成积分冻结、返还、结算及积分流水记录。","完成信用分调整、信用等级展示和双向评价。"])]
for name,items in sections:
    H(name,2)
    for x in items:B(x)
    if name.startswith("3.2"):pic(assets/"task_flow.png",15,"图 2  微任务核心业务流程")

H("3.4 治理与管理后台",2)
for x in ["完成举报提交、举报认领、证据查看、处理结果和责任人记录。","完成任务下架、用户封禁、异常任务退款和技能启停等管理操作。","完成用户、任务、举报与人工处理工作量统计。","完成趋势图、任务分类图和管理概览卡片。"]:B(x)
pic(ui/"admin_dashboard.png",15.6,"图 3  管理后台当前界面")
H("3.5 数据库、工程化与交付基础",2)
for x in ["完成核心业务数据库结构、外键约束、业务索引和增量迁移脚本。","当前设计保持 16 张业务表，覆盖用户、任务、订单、技能、积分、信用、评价、通知和举报。","完成后端 Maven 工程、前端 Vite 工程、环境变量配置和容器化部署基础。","完成 Android Capacitor 工程配置，并已生成调试 APK。","已准备后端回归、浏览器端到端、并发事务和性能验证脚本；正式测试将在开发冻结后执行。"]:B(x)
d.add_page_break()

H("四、当前界面成果")
P("以下图片由当前项目版本和本地演示数据生成，用于展示现阶段已经完成的界面与交互成果。桌面端和 Android 端共用核心业务逻辑，并针对不同视口采用响应式布局。")
H("4.1 Web 前端任务广场",2);pic(ui/"frontend_task_square.png",15.8,"图 4  Web 端任务广场")
P("任务广场已经具备搜索、分类、状态筛选、分页、任务摘要和用户入口，能够作为校园微任务的统一发现页面。")
H("4.2 Web 前端任务详情",2);pic(ui/"frontend_task_detail.png",15.8,"图 5  Web 端任务详情与操作区域")
P("任务详情页面展示任务时间、报名截止时间、积分奖励、信用要求和任务说明，并根据用户身份与任务状态显示对应操作。")
d.add_page_break()

H("4.3 Android / 移动端界面",2)
mt=d.add_table(rows=1,cols=2);mt.alignment=WD_TABLE_ALIGNMENT.CENTER
for i,(img,label) in enumerate([(ui/"android_task_square.png","移动端任务广场"),(ui/"android_profile.png","移动端个人中心")]):
    cell=mt.cell(0,i);marg(cell);p=cell.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(img),width=Cm(6.6));p=cell.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;R(p,label,True,DARK,10)
cap("图 6  Android / 移动端适配界面（390 × 844 模拟终端视口）")
P("移动端已经完成底部导航、卡片布局、筛选入口、个人中心和关键操作区域适配。Android 工程通过 Capacitor 承载同一套前端页面，目前已生成调试 APK，后续将在正式测试阶段完成多型号真机兼容性检查。")
H("五、剩余工作与正式测试安排")
H("5.1 剩余约 5% 的开发工作",2)
for x in ["统一少量页面间距、提示文案、空状态和加载状态。","复核个别边界条件下的按钮状态、异常提示和管理操作反馈。","整理生产环境配置示例、部署步骤、初始化说明和交付目录。","完成 Android 图标、启动页及不同系统版本适配细节复核。"]:B(x)
H("5.2 尚未开展的正式测试工作",2)
tt=T(["测试类型","计划内容","当前状态"],[
("功能测试","逐项验证需求、角色权限和业务状态流转","未开始"),("接口与集成测试","验证前后端接口、数据库事务和异常返回","未开始"),("兼容性测试","桌面浏览器、移动浏览器及 Android 真机","未开始"),("性能测试","任务列表、技能匹配和并发结算","未开始"),("安全性测试","认证、越权、上传、输入校验和会话失效","未开始"),("回归与验收测试","缺陷修复后执行全流程回归","未开始")],[3.6,8.2,3.1])
for row in tt.rows[1:]:
    for r in row.cells[2].paragraphs[0].runs:r.font.color.rgb=RGBColor.from_string(ORANGE)
H("5.3 主要风险与应对",2)
T(["风险","等级","应对措施"],[("需求变更","中","冻结本阶段范围，新增需求进入后续迭代"),("Android 机型差异","中","安排主流版本真机验证，优先处理阻断问题"),("集中测试发现流程缺陷","中","核心流程优先测试，预留 9 月 12 日集中修复回归"),("部署环境差异","低","使用环境变量和迁移脚本统一配置")],[4,2.2,8.8])
H("六、结论")
P("CampusLink 当前开发完成度约为 95%，已经完成核心业务、主要管理能力、Web 页面和 Android 工程适配，系统已接近开发完成。正式集中测试工作尚未开展，将在剩余开发内容冻结后按计划进行。")
p=d.add_paragraph();p.paragraph_format.first_line_indent=Cm(.74);R(p,"项目组确认：",True,DARK);R(p,"是，可以在第三阶段第 2 周（2026 年 9 月 13 日前）完成全部开发和测试工作。",True,GREEN)
for t in d.tables:
    for row in t.rows:row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
d.core_properties.title="CampusLink 项目开发进展报告";d.core_properties.subject="第三阶段第2周节点前进展说明";d.core_properties.author="CampusLink 项目组"
d.save(out);print(out)
