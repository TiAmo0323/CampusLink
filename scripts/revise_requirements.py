"""Generate revised outline §§2.4–2.6; use formal version 2 for §§1–2.3."""
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'CampusLink需求分析与详细设计书_后续章节修订版.docx'
doc=Document(ROOT/'CampusLink需求分析与详细设计书（大纲）.docx')
start=next(p._p for p in doc.paragraphs if p.text=='2.4 双终端适配与验证方案' and p.style.name.startswith('Heading'))
body=doc._element.body
for child in list(body):
    if child is start:break
    body.remove(child)
first=doc.paragraphs[0]
first.insert_paragraph_before('CampusLink需求分析与详细设计书\n后续章节修订版',style='Title')
first.insert_paragraph_before('修订日期：2026-09-05')
first.insert_paragraph_before('适用范围：本文件替代原大纲第2.4—2.5节，新增第2.6节统一规则。第1—2.3节继续采用《正式版2》，不采用旧大纲前半部分。第2.6节明确列出的冲突条款作为本轮修订决议；其他需求不变。两个原PDF及原Word均保留。')
first.insert_paragraph_before('阅读顺序：正式版2（第1—2.3节）→本文件（第2.4—2.6节）。表编号接续正式版2表2-42；验收实测结果见docs/VERIFICATION_REPORT.md，不以设计目标替代测试结论。')
def paragraphs():
 yield from doc.paragraphs
 for t in doc.tables:
  for r in t.rows:
   for c in r.cells:yield from c.paragraphs
mapping={'表2-37':'表2-43','表2-38':'表2-44','表2-39':'表2-46','表2-41':'表2-45'}
for p in paragraphs():
 text=p.text
 for a,b in mapping.items():text=text.replace(a,b).replace(a.replace('表','表 '),b)
 text=text.replace('输入手机号→获取验证码→设置密码→提交','输入用户名、密码、昵称及学号或校园邮箱→提交')
 text=text.replace('注册成功，跳转登录页，USER 表新增记录','注册成功，签发登录凭证进入首页；user新增记录并生成初始积分流水')
 text=text.replace('返回 Top10 推荐列表，按 MatchScore 降序排列','默认10条分页；MatchScore降序，同分互补优先，再按用户ID排序')
 text=text.replace('任务标记超时，接取者信用-2，发布者可重新选择','任务和订单转ABNORMAL，原因为TIMEOUT；接取者实际信用扣分受0分边界限制；重新选择须设置新期限并保留分配历史')
 text=text.replace('接取者信用+5','接取者信用按+5规则增加（最高100，记录实际变化）')
 if 'P1 用例作为演示加分项' in text:text='TC-01—TC-06均为第一版P0必测项。另执行登录安全、撤回、阶段取消、驳回上限、重新分配、技能双确认、分页及移动端回归。'
 if '该示例包含 MySQL 数据库、Redis' in text:text='第一版部署包含MySQL、Spring Boot后端、Nginx前端三个服务；Redis不是必需依赖。刷新会话保存在MySQL，当前不引入无实际业务用途的Redis。'
 if 'image: redis:7-alpine' in text:text=(ROOT/'docker-compose.yml').read_text(encoding='utf-8')
 if 'MySQL 与 Redis 数据通过 Docker Volume' in text:text='配置.env后执行docker compose up -d --build。MySQL和上传文件分别使用持久卷。迁移使用Flyway；不通过清空数据库升级。生产环境禁用演示账号并配置随机令牌密钥。'
 if text!=p.text:p.text=text
for t in doc.tables:
 if any('TC-01' in c.text for r in t.rows for c in r.cells):
  for r in t.rows[1:]:r.cells[-1].text='P0'
doc.add_heading('2.6 跨章节冲突修正与补充设计',level=2)
doc.add_paragraph('下列为明确的需求修订，不以“代码如此实现”为默认依据。涉及正式版2的覆盖条款：表2-41的积分/信用非零约束、仅注册登录可匿名的接口总则、待执行取消双方扣分歧义、订单唯一性与重新选择历史、技能互助完成责任及新增字段。未列明的正式版2内容继续有效。')
for line in (ROOT/'docs/REVISION_BASELINE.md').read_text(encoding='utf-8').splitlines():
 if line.startswith('- '):doc.add_paragraph(line[2:])
doc.add_paragraph('补充安全约束：头像JPG/PNG≤5MB；完成凭证JPG/PNG/PDF≤10MB。服务端检查大小、扩展名及文件内容；前端PDF采用链接查看，图片采用自适应预览。PDF基本类型校验不等同恶意内容检测，生产部署仍需文件扫描或隔离下载策略。')
doc.add_heading('2.6.1 数据库增补与兼容迁移',level=3)
doc.add_paragraph('保留14张业务主表，新增refresh_session与order_assignment_history两张技术/审计表，共16张应用表（Flyway表另计）。V1—V4不改写；增量V5迁移保存历史。新流水必须写冻结余额，旧流水无法可靠反推时保留NULL；历史完成记录允许submitter_id为NULL，新提交必须写入。')
rows=[
('user','last_login_at；failed_login_attempts INT DEFAULT 0；locked_until DATETIME NULL；token_version INT DEFAULT 0','最近登录、失败锁定、旧凭证撤销；last_login_at已存在'),
('user_skill','available_time VARCHAR(255) NULL','可提供时间，不与available_mode混用'),
('task','estimated_duration_minutes INT DEFAULT 60；abnormal_reason VARCHAR(50)；cancel_reason VARCHAR(255)；is_delisted BOOLEAN DEFAULT FALSE；delist_reason VARCHAR(255)','预计时长与独立下架；version沿用'),
('task_order','assignment_round INT DEFAULT 1；due_at DATETIME NULL','同一订单分轮履约；version沿用'),
('task_completion','submitter_id BIGINT NULL FK user.id；assignment_round INT DEFAULT 1','历史归属与轮次'),
('skill_exchange','provider_completed_at DATETIME NULL','提供方先完成，发起方再确认'),
('point_transaction','before_frozen INT NULL；after_frozen INT NULL','冻结余额变化可审计'),
('report','previous_status VARCHAR(30) NULL','异常恢复前态；ORDER关联订单，该字段已存在'),
('refresh_session','id BIGINT PK；user_id BIGINT FK；token_hash VARCHAR(64) UNIQUE；expires_at/created_at DATETIME','数据库仅存刷新凭证SHA-256哈希；轮换删除旧记录'),
('order_assignment_history','id BIGINT PK；order_id BIGINT FK；previous_accepter_id/new_accepter_id BIGINT FK；previous_round INT；previous_started_at/previous_due_at DATETIME NULL；reason VARCHAR(255)；created_at DATETIME','完整保留前轮接取者和期限；不删除旧凭证')]
t=doc.add_table(rows=1,cols=3);t.style='Table Grid'
for c,v in zip(t.rows[0].cells,['表','字段','用途/兼容规则']):c.text=v
for row in rows:
 for c,v in zip(t.add_row().cells,row):c.text=v
doc.add_paragraph('约束修正：change_value=after_score-before_score，允许0；change_amount=after_balance-before_balance，允许0但新事件至少一种余额发生变化。余额/信用范围检查保留。所有核心外键RESTRICT；通知删除策略沿用现有数据库约束，当前无物理删除用户接口。')
doc.add_paragraph('并发策略：既有履约操作按任务→订单→用户ID升序锁定；发布新任务先锁发布者用户，再插入任务/冻结积分，避免外键共享锁升级死锁；评价先按用户ID升序锁定双方后插入。每个超时处理单独事务；状态条件与更新行数必须校验。')
doc.add_heading('2.6.2 状态和接口约定',level=3)
doc.add_paragraph('Java核心状态集中于domain/States.java；任务与订单共用Task状态，申请、完成记录、互助、举报分别使用对应枚举值。前端映射同一字符串协议。统一响应{code,message,data}，成功code=200；失败code与HTTP状态400/401/403/404/429/500对齐，替代正式版2表2-23的40101等编码。公开读取/api/public无需登录；刷新/退出通过刷新凭证验证，其他业务鉴权。')
doc.add_paragraph('集合统一返回{page,size,total,records}，total是完整结果数量；聚合详情中的完成记录附completionTotal；技能档案分别返回offersTotal/needsTotal并分页，不能以截断代替分页。')
for s in [
'POST /api/auth/refresh、/api/auth/logout：刷新凭证轮换与退出；退出撤销该用户所有会话，旧访问令牌不恢复。',
'POST /api/applications/{id}/withdraw：只撤回本人PENDING申请；GET /api/tasks/mine?kind=published|accepted|applied&page=1&size=10。',
'POST /api/tasks/{id}/cancel：待执行阶段reason必填；POST /api/tasks/{id}/reassign：applicationId、dueAt、reason必填。',
'GET /api/orders/{id}?page=1&size=10：当前双方完整详情，原接取者只读本人历史；新接取者不得操作原轮次。',
'POST /api/exchanges/{id}/accept|reject|start|complete|cancel；complete按调用角色先记录提供方时间、再由发起方完成。',
'POST /api/reports；POST /api/admin/reports/{id}/claim；POST /api/admin/reports/{id}/handle；POST /api/admin/tasks/{id}/delist。',
'POST /api/files/images为头像接口；POST /api/files/proofs为完成凭证接口。'
]:doc.add_paragraph(s)
doc.add_heading('2.6.3 七项开发目标与验收边界',level=3)
for line in (ROOT/'docs/REVISION_BASELINE.md').read_text(encoding='utf-8').splitlines():
 if len(line)>2 and line[0].isdigit() and line[1:3]=='. ':doc.add_paragraph(line)
doc.add_paragraph('浏览器设备模拟验证不等同真机Android运行。移动验收须注明视口、浏览器、是否实际在Android WebView执行；APK编译成功不等同真机功能通过。压测必须记录首次失败和修复后的重测，不得只报告平均耗时或掩盖错误请求。')
doc.core_properties.title='CampusLink需求分析与详细设计书：后续章节修订版'
doc.save(OUT)
print(OUT)
if __name__=='__main__':
 import win32com.client
 word=win32com.client.DispatchEx('Word.Application');word.Visible=False;word.DisplayAlerts=0
 try:
  opened=word.Documents.Open(str(OUT));opened.Fields.Update();opened.Save();opened.ExportAsFixedFormat(str(OUT.with_suffix('.pdf')),17);opened.Close(False);print(OUT.with_suffix('.pdf'))
 finally:word.Quit()
