# -*- coding: utf-8 -*-
from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
from matplotlib import font_manager

font_manager.fontManager.addfont(r"C:\Windows\Fonts\msyh.ttc")
plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "DejaVu Sans"]
plt.rcParams["axes.unicode_minus"] = False
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"D:\CampusLink")
SOURCE = ROOT / "CampusLink需求分析与详细设计书（大纲）.docx"
OUTPUT = ROOT / "CampusLink软件系统分析设计与实现报告.docx"
ASSETS = ROOT / "report_assets"

GREEN = "2F6558"
LIGHT_GREEN = "E9F1ED"
LIGHT_GRAY = "F3F5F7"
GOLD = "B97824"
TEXT = "25342F"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def configure_run(run, size: float = 10.5, bold: bool = False, color: str = TEXT, font: str = "宋体") -> None:
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_body(doc: Document, text: str = "", *, bold: bool = False, indent: bool = True,
             align=WD_ALIGN_PARAGRAPH.JUSTIFY) -> None:
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(5)
    if indent:
        p.paragraph_format.first_line_indent = Pt(21)
    r = p.add_run(text)
    configure_run(r, 10.5, bold)


def add_bullets(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.74)
        p.paragraph_format.first_line_indent = Cm(-0.5)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.25
        configure_run(p.add_run(item), 10.5)


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_paragraph(text, style=f"Heading {level}")
    set_keep_with_next(p)


def add_table(doc: Document, headers: Sequence[str], rows: Sequence[Sequence[object]], widths=None,
              font_size: float = 9.0) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, GREEN)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        configure_run(p.add_run(str(text)), font_size, True, "FFFFFF", "微软雅黑")
        if widths:
            cell.width = Cm(widths[i])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            set_cell_margins(cell)
            if row_index % 2 == 1:
                set_cell_shading(cell, LIGHT_GRAY)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            configure_run(p.add_run(str(value)), font_size, False, TEXT)
            if widths:
                cell.width = Cm(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    configure_run(p.add_run(text), 9, False, "5E6D68")


def add_code(doc: Document, lines: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.right_indent = Cm(0.6)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.0
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F4F6F5")
    p_pr.append(shd)
    r = p.add_run(lines)
    configure_run(r, 8.5, False, "2E3A36", "Consolas")


def page_break(doc: Document) -> None:
    doc.add_page_break()


def fix_source_content(doc: Document) -> None:
    replacements = {
        "需求分析与详细设计书": "软件系统分析、设计与实现报告",
        "账户与个人中心的入口设计以登录/注册/找回密码三个 Tab 整合为同一卡片，PC 端采用居中卡片布局配合图形验证码与短信验证码双重校验，移动端精简为手机号+短信验证码一键登录，覆盖校园场景下手机优先的访问习惯。具体界面设计如图2-11所示。":
            "账户与个人中心入口采用登录、注册独立页面与统一卡片式布局。当前版本使用用户名或邮箱加密码完成登录，Web 与 Android 端复用同一 Vue 页面；移动端通过响应式布局压缩留白并放大触控区域。短信验证码、找回密码属于后续扩展范围。具体界面设计如图2-11所示。",
        "为体现系统可直接部署的工程完整性，给出基于 Docker Compose 的最小部署示例。该示例包含 MySQL 数据库、Redis 缓存、SpringBoot 后端与 Nginx 前端代理四个容器，通过 docker-compose.yml 一键启动。":
            "为体现系统可直接部署的工程完整性，当前仓库给出 Docker Compose 部署方案。编排包含 MySQL 8、Spring Boot 后端与 Nginx 前端三个服务；当前业务规模不依赖 Redis，状态与事务一致性由数据库、服务层事务及条件更新共同保障。",
        "启动命令为 docker-compose up -d，系统通过 Nginx 80 端口对外提供前端静态资源与后端 API 反向代理；MySQL 与 Redis 数据通过 Docker Volume 持久化，避免容器重建导致数据丢失。":
            "启动命令为 docker compose up --build -d。Nginx 的 80 端口提供前端资源并代理 API；MySQL 数据和用户上传文件分别使用 campuslink_mysql、campuslink_uploads 命名卷持久化，避免容器重建导致数据丢失。"
    }
    for p in doc.paragraphs:
        if p.text in replacements:
            p.text = replacements[p.text]
    cover = doc.tables[0]
    cover.cell(4, 1).text = "2026 年 9 月"
    # TC-01 位于原报告最后一张测试表。
    for table in doc.tables:
        for row in table.rows:
            if row.cells and row.cells[0].text.strip() == "TC-01":
                values = ["TC-01", "用户注册", "数据库可访问", "输入用户名、密码、昵称、学号与邮箱后提交",
                          "注册成功并跳转登录页，USER 表新增记录，初始化积分与信用", "P0"]
                for c, v in zip(row.cells, values):
                    c.text = v


def draw_architecture(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(11, 5.8), dpi=180)
    ax.set_xlim(0, 11)
    ax.set_ylim(0, 6)
    ax.axis("off")
    colors = ["#E9F1ED", "#E8EFF7", "#FFF2DF", "#F3EEF7"]
    layers = [
        (4.8, "终端层", "响应式 Web（浏览器）  ｜  Capacitor Android（Pixel 8 / API 36）"),
        (3.45, "表示层", "Vue 3 + Vite + Element Plus + Pinia + Vue Router + Axios"),
        (2.1, "业务与接口层", "Spring Boot MVC  ｜  Controller  ｜  Service  ｜  鉴权/校验/异常处理"),
        (0.75, "数据与基础设施层", "MyBatis-Plus + Flyway  ｜  MySQL 8 / H2  ｜  文件存储  ｜  Nginx")
    ]
    for i, (y, title, detail) in enumerate(layers):
        box = FancyBboxPatch((0.7, y), 9.6, 0.85, boxstyle="round,pad=0.03,rounding_size=0.08",
                             facecolor=colors[i], edgecolor="#2F6558", linewidth=1.3)
        ax.add_patch(box)
        ax.text(1.0, y + 0.55, title, fontsize=13, weight="bold", color="#2F6558", va="center")
        ax.text(2.65, y + 0.42, detail, fontsize=10.5, color="#25342F", va="center")
        if i < len(layers) - 1:
            ax.add_patch(FancyArrowPatch((5.5, y - 0.05), (5.5, y - 0.43), arrowstyle="-|>",
                                         mutation_scale=13, color="#76867F", linewidth=1.1))
    ax.text(5.5, 5.8, "CampusLink 当前实现总体架构", ha="center", fontsize=16, weight="bold", color="#234F45")
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def draw_task_flow(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(12, 5.4), dpi=180)
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 5.5)
    ax.axis("off")
    nodes = [
        (0.5, 3.6, "招募中\nRECRUITING"), (2.7, 3.6, "待执行\nWAIT_EXECUTE"),
        (4.9, 3.6, "进行中\nIN_PROGRESS"), (7.1, 3.6, "待验收\nWAIT_ACCEPTANCE"),
        (9.3, 3.6, "已完成\nCOMPLETED"), (0.5, 1.2, "已取消\nCANCELLED"),
        (6.0, 1.2, "异常中\nABNORMAL")
    ]
    for x, y, label in nodes:
        box = FancyBboxPatch((x, y), 1.7, 0.75, boxstyle="round,pad=0.04,rounding_size=0.08",
                             facecolor="#E9F1ED" if "异常" not in label else "#FFF0E5",
                             edgecolor="#2F6558" if "异常" not in label else "#B97824", linewidth=1.4)
        ax.add_patch(box)
        ax.text(x + 0.85, y + 0.38, label, ha="center", va="center", fontsize=9.5, color="#25342F")
    def arrow(x1, y1, x2, y2, label=""):
        ax.add_patch(FancyArrowPatch((x1, y1), (x2, y2), arrowstyle="-|>", mutation_scale=13,
                                     color="#60736B", linewidth=1.2, connectionstyle="arc3,rad=0"))
        if label:
            ax.text((x1+x2)/2, (y1+y2)/2 + 0.12, label, ha="center", fontsize=8.5, color="#4A5A54")
    arrow(2.2, 3.98, 2.7, 3.98, "确认申请")
    arrow(4.4, 3.98, 4.9, 3.98, "开始")
    arrow(6.6, 3.98, 7.1, 3.98, "提交")
    arrow(8.8, 3.98, 9.3, 3.98, "验收")
    arrow(7.1, 3.65, 6.6, 3.65, "驳回重提")
    arrow(1.35, 3.55, 1.35, 1.95, "招募期取消")
    arrow(5.75, 3.55, 6.7, 1.95, "争议/超时")
    arrow(7.95, 3.55, 7.4, 1.95, "争议")
    ax.text(6.85, 0.72, "管理员裁决：退款 REFUND / 结算 SETTLE / 恢复 RESUME", ha="center",
            fontsize=9.5, color="#874F13")
    ax.text(6, 5.05, "微任务业务状态流转", ha="center", fontsize=16, weight="bold", color="#234F45")
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def draw_deployment(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(11, 5.7), dpi=180)
    ax.set_xlim(0, 11)
    ax.set_ylim(0, 6)
    ax.axis("off")
    boxes = [
        (0.5, 4.3, 2.4, 0.9, "浏览器 / Android", "HTTP(S)"),
        (4.0, 4.3, 2.5, 0.9, "Nginx 前端", "静态资源 + /api 代理"),
        (7.8, 4.3, 2.5, 0.9, "Spring Boot", "8080 / 上传目录"),
        (7.8, 1.8, 2.5, 0.9, "MySQL 8", "campuslink_mysql 卷"),
        (4.0, 1.8, 2.5, 0.9, "文件持久化", "campuslink_uploads 卷")
    ]
    for x, y, w, h, title, detail in boxes:
        box = FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.05,rounding_size=0.08",
                             facecolor="#F4F8F6", edgecolor="#2F6558", linewidth=1.3)
        ax.add_patch(box)
        ax.text(x+w/2, y+0.59, title, ha="center", fontsize=11, weight="bold", color="#2F6558")
        ax.text(x+w/2, y+0.27, detail, ha="center", fontsize=8.8, color="#53635D")
    def ar(a, b):
        ax.add_patch(FancyArrowPatch(a, b, arrowstyle="<|-|>", mutation_scale=12, color="#60736B", linewidth=1.2))
    ar((2.9, 4.75), (4.0, 4.75))
    ar((6.5, 4.75), (7.8, 4.75))
    ar((9.05, 4.25), (9.05, 2.7))
    ar((7.8, 4.48), (6.5, 2.5))
    ax.text(5.5, 5.65, "Docker Compose 三服务部署拓扑", ha="center", fontsize=16, weight="bold", color="#234F45")
    ax.text(5.5, 0.75, "开发环境可不启动 Docker：Spring Boot 默认使用 H2 内存数据库；Android 模拟器访问宿主机地址 10.0.2.2。",
            ha="center", fontsize=9.2, color="#4A5A54")
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def add_picture(doc: Document, path: Path, width_inches: float, caption: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width_inches))
    add_caption(doc, caption)


def build_report() -> None:
    ASSETS.mkdir(exist_ok=True)
    architecture = ASSETS / "architecture.png"
    task_flow = ASSETS / "task_flow.png"
    deployment = ASSETS / "deployment.png"
    draw_architecture(architecture)
    draw_task_flow(task_flow)
    draw_deployment(deployment)

    doc = Document(SOURCE)
    fix_source_content(doc)

    # 保持原文档页面设置，新增内容单独起页。
    page_break(doc)
    add_heading(doc, "3 系统实现", 1)
    add_heading(doc, "3.1 文档基线与实现范围", 2)
    add_body(doc, "本章及后续章节以 2026 年 9 月 4 日工作区 D:\\CampusLink、Git 提交 6a9c947 为实现基线。第 1～2 章给出需求与详细设计，本章起记录已完成代码、构建配置和实际验收结果；当早期原型示例与本章存在差异时，以当前实现和可执行配置为准。")
    add_table(doc, ["项目", "当前基线"], [
        ("系统名称", "CampusLink——基于信用积分与技能匹配的校园微任务互助平台"),
        ("产品形态", "响应式 Web 应用 + Capacitor Android 应用 + Java REST 服务"),
        ("版本", "1.0.0（课程设计可运行版本）"),
        ("源码基线", "Git 提交 6a9c947，主分支 main"),
        ("数据库", "部署使用 MySQL 8；快速演示和自动化测试使用 H2 2.x 的 MySQL 兼容模式"),
        ("实现边界", "不含短信验证码、找回密码、实时聊天、地图定位、真实支付、WebSocket、Redis 与大模型推荐"),
    ], [4.0, 12.0])

    add_heading(doc, "3.2 技术选型与运行环境", 2)
    add_table(doc, ["层次/模块", "技术与版本", "用途"], [
        ("后端运行时", "Java 21、Spring Boot 3.4.5", "REST API、依赖注入、事务、定时任务与统一异常处理"),
        ("数据访问", "MyBatis-Plus、Flyway", "实体映射、条件更新、分页查询和数据库版本迁移"),
        ("接口文档", "SpringDoc OpenAPI 2.8.8", "生成 /swagger-ui.html 与 /v3/api-docs"),
        ("数据库", "MySQL 8 / H2 MySQL 模式", "持久化部署 / 本地快速演示与集成测试"),
        ("Web 前端", "Vue 3.5、Vite 6.2、Pinia 3、Vue Router 4.5", "组件化页面、状态管理、路由与构建"),
        ("UI 与通信", "Element Plus 2.9、Axios 1.8", "界面组件和 HTTP 请求"),
        ("Android", "Capacitor 8.4.3、SDK 36、Gradle 8.14.3", "复用 Web 业务界面并封装为 Android 原生工程"),
        ("部署", "Docker Compose、Nginx", "MySQL/后端/前端三服务编排和反向代理"),
    ], [3.3, 5.3, 7.4])
    add_picture(doc, architecture, 6.7, "图 3-1  当前实现总体架构")

    add_heading(doc, "3.3 后端实现", 2)
    add_body(doc, "后端采用 Controller—Service—Mapper—Domain 分层。Controller 负责参数绑定与接口边界；Service 负责事务、权限和状态机；Mapper 基于 MyBatis-Plus 访问数据；Domain 与 DTO 分离，避免直接以持久化实体接收外部输入。全局异常处理器将参数错误、业务错误和系统错误统一转换为 ApiResponse。")
    add_table(doc, ["包/组件", "代表类", "职责"], [
        ("controller", "AuthController、TaskController、SkillController、GovernanceController", "暴露账户、任务、技能、治理及文件 REST 接口"),
        ("service", "TaskService、PointCreditService、AbnormalTaskService、TaskTimeoutService", "业务状态流转、积分信用结算、异常裁决与超时扫描"),
        ("config", "AuthInterceptor、TokenService、WebConfig、BootstrapData", "令牌认证、请求拦截、跨域/静态资源、演示数据"),
        ("mapper", "14 个 Mapper 接口", "对应业务表的增删改查、条件更新与统计"),
        ("domain/dto", "CampusTask、TaskOrder、Requests", "持久化模型与经过校验的请求模型"),
        ("common", "ApiResponse、BusinessException、GlobalExceptionHandler", "统一响应码、业务异常和异常映射"),
    ], [3.0, 6.0, 7.0])
    add_heading(doc, "3.3.1 鉴权与账户实现", 3)
    add_body(doc, "注册时校验用户名、学号和邮箱唯一性，使用 BCrypt 保存密码摘要，并初始化 200 可用积分和信用分。登录支持用户名或邮箱，校验成功后生成带 HMAC 签名、默认有效期 24 小时的令牌。除公共接口外，请求必须携带 Authorization: Bearer <token>；拦截器解析令牌并写入 AuthContext，管理员接口再次校验 ADMIN 角色。")
    add_heading(doc, "3.3.2 微任务与积分事务", 3)
    add_body(doc, "发布任务在同一事务中检查余额、冻结奖励积分并写入任务；确认申请使用任务状态和版本条件更新，确保并发点击时仅生成一个履约订单；验收通过后从发布者冻结积分结算至接取者可用积分，同时写入不可篡改的积分流水。招募期取消执行退款；履约期争议进入异常任务流程，由管理员选择退款、结算或恢复。")
    add_picture(doc, task_flow, 7.0, "图 3-2  微任务状态流转与异常分支")
    add_heading(doc, "3.3.3 技能匹配实现", 3)
    add_body(doc, "系统将用户技能需求与候选人的技能供给取交集，过滤无匹配技能的候选人，并按技能覆盖度 50%、信用分 20%、历史评价 20%、近 90 日活跃度 10% 计算综合分。结果最多返回 10 条，同时标记对方需求是否与本人供给形成双向互补，保证推荐结果可解释。")
    add_heading(doc, "3.3.4 超时与异常治理", 3)
    add_body(doc, "定时任务默认每 10 分钟扫描 WAIT_EXECUTE、IN_PROGRESS 订单。超过任务时间与 24 小时宽限期后，以状态条件更新方式将订单置为 ABNORMAL，仅首次更新者写入异常举报、扣减信用并发送通知，从而保证重复扫描幂等。管理员裁决同样采用 PROCESSING→HANDLING 的原子抢占，避免同一异常被重复结算。")

    add_heading(doc, "3.4 前端实现", 2)
    add_body(doc, "前端采用单页应用结构。App.vue 提供全局品牌栏、主导航、身份展示和上下文返回按钮；Router 根据页面元数据执行登录与管理员权限守卫；Pinia 保存用户会话；Axios 请求拦截器附加令牌，响应拦截器统一处理业务错误和 401 失效。页面通过 CSS 媒体查询适配桌面与移动宽度。")
    add_table(doc, ["路由", "页面", "访问要求", "主要能力"], [
        ("/login", "登录", "游客", "用户名/邮箱与密码登录"),
        ("/register", "注册", "游客", "创建学生账户"),
        ("/tasks", "任务广场", "公开", "分页、关键词、分类和状态筛选"),
        ("/tasks/new", "发布任务", "登录", "录入任务信息并冻结积分"),
        ("/tasks/:id", "任务详情", "公开", "查看、申请、取消和申请人管理"),
        ("/mine", "我的任务", "登录", "区分我发布与我接取的任务"),
        ("/orders/:id", "履约订单", "登录且为参与者", "开始、提交、验收、驳回和异常"),
        ("/skills", "技能互助", "登录", "维护供需、查看匹配和处理互助申请"),
        ("/wallet", "积分信用", "登录", "余额、冻结积分、积分与信用流水"),
        ("/messages", "消息中心", "登录", "通知分页与已读"),
        ("/profile", "个人中心", "登录", "资料编辑与账户信息"),
        ("/admin", "管理台", "管理员", "统计、用户、任务、举报、异常和技能治理"),
    ], [3.0, 2.5, 3.0, 7.0], 8.5)

    add_heading(doc, "3.5 Android 移动端实现", 2)
    add_body(doc, "Android 工程由 Capacitor 8.4.3 生成，应用标识为 com.campuslink.app，最低 API 24，编译和目标 API 均为 36。原生容器加载 Vite 构建产物，移动端使用 Hash 路由避免 WebView 深链接刷新错误。调试环境通过 VITE_API_BASE=http://10.0.2.2:8080 访问模拟器宿主机，并仅在 debug Manifest 中允许明文 HTTP；正式发布必须切换 HTTPS。")
    add_table(doc, ["适配点", "实现方式", "验收依据"], [
        ("路由", "原生平台 createWebHashHistory，Web 使用 createWebHistory", "应用内页面跳转正常"),
        ("网络", "Android 模拟器使用 10.0.2.2，Capacitor scheme 为 http", "Pixel 8 模拟器登录与业务请求成功"),
        ("安全边界", "cleartext 仅用于本地调试，release 主清单不开放", "Debug 可连接本机，生产说明要求 HTTPS"),
        ("布局", "响应式导航、卡片重排、触控尺寸与安全区", "1080×2400 Pixel 8 画面可用"),
        ("返回导航", "任务详情、发布页、订单页提供语义化左箭头；系统返回键仍可用", "核心页面可返回任务广场或我的任务"),
        ("构建", "npm run android:sync + gradlew assembleDebug", "生成 app-debug.apk"),
    ], [3.0, 7.0, 6.0])

    page_break(doc)
    add_heading(doc, "4 接口设计与实现", 1)
    add_heading(doc, "4.1 接口约定", 2)
    add_body(doc, "服务默认监听 8080 端口，采用 JSON REST 接口。成功响应格式为 {\"code\":0,\"message\":\"success\",\"data\":...}；失败时 code 使用 400、401、403、404 或 500 等语义码，message 返回可展示的原因。图片上传使用 multipart/form-data，字段名为 file，单文件和单请求上限均为 10 MB。")
    add_code(doc, 'Authorization: Bearer <登录接口返回的 token>\nContent-Type: application/json\n\n{\n  "code": 0,\n  "message": "success",\n  "data": { ... }\n}')

    api_rows = [
        ("POST", "/api/auth/register", "公开", "注册并初始化积分/信用"),
        ("POST", "/api/auth/login", "公开", "登录并获取令牌"),
        ("GET", "/api/public/tasks", "公开", "分页与条件查询任务"),
        ("GET", "/api/public/tasks/{id}", "公开", "任务详情"),
        ("GET", "/api/public/skills", "公开", "启用的技能分类"),
        ("GET", "/api/users/me", "用户", "本人资料"),
        ("PUT", "/api/users/me", "用户", "更新本人资料"),
        ("GET", "/api/wallet", "用户", "余额及积分/信用流水"),
        ("GET", "/api/notifications", "用户", "分页通知"),
        ("POST", "/api/notifications/{id}/read", "用户", "通知标记已读"),
        ("POST", "/api/files/images", "用户", "上传图片凭证"),
        ("POST", "/api/tasks", "用户", "发布任务并冻结积分"),
        ("GET", "/api/tasks/mine", "用户", "我发布/接取的任务"),
        ("POST", "/api/tasks/{id}/applications", "用户", "申请接取任务"),
        ("GET", "/api/tasks/{id}/applications", "发布者", "查看任务申请列表"),
        ("POST", "/api/applications/{id}/accept", "发布者", "原子确认接取者"),
        ("POST", "/api/tasks/{id}/cancel", "发布者", "招募期取消并退款"),
        ("POST", "/api/tasks/{id}/abnormal", "任务双方", "发起任务异常"),
        ("GET", "/api/orders/{id}", "任务双方", "履约订单详情"),
        ("POST", "/api/orders/{id}/start", "接取者", "开始执行"),
        ("POST", "/api/orders/{id}/completions", "接取者", "提交结果及凭证"),
        ("POST", "/api/orders/{id}/approve", "发布者", "验收并结算"),
        ("POST", "/api/orders/{id}/reject", "发布者", "驳回并要求重提"),
        ("GET", "/api/skills/profile", "用户", "本人技能供需"),
        ("POST", "/api/skills/offers", "用户", "新增或更新技能供给"),
        ("DELETE", "/api/skills/offers/{id}", "用户", "删除技能供给"),
        ("POST", "/api/skills/needs", "用户", "新增或更新技能需求"),
        ("DELETE", "/api/skills/needs/{id}", "用户", "删除技能需求"),
        ("GET", "/api/skills/matches", "用户", "Top 10 可解释匹配"),
        ("POST", "/api/exchanges", "用户", "发起技能互助"),
        ("GET", "/api/exchanges", "用户", "查询本人互助"),
        ("POST", "/api/exchanges/{id}/{action}", "参与者", "accept/reject/start/complete"),
        ("POST", "/api/reviews", "已完成业务双方", "提交评价"),
        ("GET", "/api/reviews/users/{id}", "用户", "查询用户收到的评价"),
        ("POST", "/api/reports", "用户", "提交普通举报"),
        ("GET", "/api/reports/mine", "用户", "查询本人举报"),
        ("GET", "/api/admin/dashboard", "管理员", "聚合统计"),
        ("GET", "/api/admin/users", "管理员", "分页用户列表"),
        ("PUT", "/api/admin/users/{id}/status/{status}", "管理员", "启用/禁用账户"),
        ("GET", "/api/admin/tasks", "管理员", "分页任务列表"),
        ("GET", "/api/admin/reports", "管理员", "分页举报列表"),
        ("POST", "/api/admin/reports/{id}/handle", "管理员", "处理普通举报"),
        ("POST", "/api/admin/abnormal-tasks/{reportId}/resolve", "管理员", "异常任务裁决"),
        ("GET", "/api/admin/skills", "管理员", "全部技能分类"),
        ("POST", "/api/admin/skills", "管理员", "新增技能分类"),
        ("PUT", "/api/admin/skills/{id}/status/{status}", "管理员", "启用/停用技能分类"),
    ]
    add_heading(doc, "4.2 API 清单", 2)
    add_table(doc, ["方法", "路径", "权限", "功能"], api_rows, [1.6, 7.3, 2.8, 4.5], 8.0)
    add_heading(doc, "4.3 关键请求数据约束", 2)
    add_table(doc, ["对象", "主要字段", "校验规则"], [
        ("Register", "username、password、nickname、studentNo、email", "用户名 3～32 位；密码至少 6 位；用户名/学号/邮箱唯一"),
        ("TaskCreate", "title、category、description、location、taskTime、applicationDeadline、rewardPoints、minCreditScore", "标题≤100；奖励>0；信用 0～100；时间关系合法"),
        ("Completion", "description、proofUrl", "完成说明非空；凭证 URL≤255"),
        ("ReviewCreate", "businessType、businessId、revieweeId、rating、content", "仅完成业务参与者；星级 1～5；同一评价人每业务一次"),
        ("AbnormalCreate", "reasonType、description", "类型为 NO_CONTACT/DISPUTE/TIMEOUT/OTHER"),
        ("AbnormalResolve", "decision、result、liableUserId、creditDelta", "决策 REFUND/SETTLE/RESUME；责任人必须为任务参与者"),
    ], [3.0, 7.8, 5.3], 8.5)

    page_break(doc)
    add_heading(doc, "5 数据库实现", 1)
    add_heading(doc, "5.1 数据持久化策略", 2)
    add_body(doc, "数据库结构由 Flyway 管理。V1 创建 14 张业务表和基础索引，V2 增加异常任务恢复所需的 previous_status，V3 补充业务外键，V4 增加积分、信用、奖励和星级检查约束。已执行迁移不可回改，后续变更必须新增更高版本脚本。MySQL 命名卷保证部署数据持久化；H2 内存模式用于快速演示，进程停止后数据清空。")
    db_rows = [
        ("user", "用户账户", "身份、角色、密码摘要、积分余额、冻结积分、信用和状态"),
        ("skill", "技能分类", "技能名称、分类和启停状态"),
        ("user_skill", "技能供给", "用户可提供的技能、熟练度、说明和方式"),
        ("user_skill_need", "技能需求", "用户所需技能、优先级、说明和偏好方式"),
        ("task", "微任务", "发布者、任务内容、时空信息、奖励、信用门槛、状态和版本"),
        ("task_application", "任务申请", "任务、申请人、留言、预计完成时间和申请状态"),
        ("task_order", "履约订单", "唯一任务、双方用户、状态、时间戳和乐观版本"),
        ("task_completion", "完成提交", "订单、完成说明、凭证、审核状态和驳回原因"),
        ("skill_exchange", "技能互助", "请求人、提供者、双方技能、约定时间和状态"),
        ("point_transaction", "积分流水", "业务来源、变动额、变动前后余额、时间和备注"),
        ("credit_record", "信用流水", "业务来源、变化值、前后信用分与原因"),
        ("review", "评价", "业务类型、业务编号、评价双方、星级与内容"),
        ("report", "举报/异常", "目标、原因、原状态、处理状态、处理人和结果"),
        ("notification", "站内通知", "用户、类型、标题、内容、关联业务与已读状态"),
    ]
    add_heading(doc, "5.2 业务表数据字典", 2)
    add_table(doc, ["表名", "中文名称", "核心数据"], db_rows, [4.0, 3.0, 9.0], 8.5)
    add_heading(doc, "5.3 约束与一致性", 2)
    add_table(doc, ["约束类型", "实现", "保证"], [
        ("唯一约束", "用户名、学号、邮箱、技能名、用户技能组合、任务申请组合、任务订单、评价组合", "防止重复身份、重复申请、重复订单和重复评价"),
        ("检查约束", "信用 0～100、可用/冻结积分非负、任务奖励为正、最低信用 0～100、评价 1～5", "阻断越界业务数据"),
        ("外键", "V3 为任务、申请、订单、技能、流水、评价、举报和通知补充关联", "保证引用记录存在"),
        ("索引", "任务状态/截止时间、发布者/状态、用户流水时间、通知用户/已读/时间", "支持任务筛选和个人分页查询"),
        ("事务", "发布冻结、确认申请、验收结算、异常裁决均使用 @Transactional", "业务写入要么全部成功，要么全部回滚"),
        ("条件更新", "任务/订单状态与 version、举报处理状态作为更新条件", "防止并发重复确认与重复结算"),
    ], [3.0, 7.4, 5.6], 8.5)

    page_break(doc)
    add_heading(doc, "6 测试与验收", 1)
    add_heading(doc, "6.1 测试策略与环境", 2)
    add_body(doc, "测试覆盖后端业务闭环、前端生产构建、Android 资源同步与 Debug APK 构建，并在 Pixel 8 / Android 16（API 36）模拟器上完成应用启动和登录验证。后端集成测试每次通过 Flyway clean/migrate 重建 H2 数据库并加载演示数据，避免用例间污染。")
    add_table(doc, ["层次", "工具/环境", "验收内容"], [
        ("后端", "Maven、Spring Boot Test、MockMvc、H2", "接口权限、状态机、积分信用、并发、文件上传"),
        ("Web", "Node.js、npm、Vite", "依赖安装、Vue 模板编译和生产资源构建"),
        ("Android", "Android Studio Quail 3、JVM 21、Gradle 8.14.3", "Capacitor 同步、assembleDebug、应用启动"),
        ("模拟器", "Pixel 8、Android 16 / API 36、x86_64", "首页、后端连接、登录和响应式界面"),
        ("部署", "Docker Compose 配置", "三服务依赖、端口、环境变量和持久化卷"),
    ], [3.0, 6.0, 7.0])
    add_heading(doc, "6.2 自动化集成测试结果", 2)
    add_table(doc, ["编号", "测试场景", "关键断言", "结果"], [
        ("IT-01", "完整任务闭环", "发布冻结→申请→确认→开始→提交→验收→积分结算→评价信用联动", "通过"),
        ("IT-02", "技能匹配与管理员鉴权", "匹配接口可用，普通用户不能访问管理接口", "通过"),
        ("IT-03", "异常退款裁决", "异常任务被管理员退款，积分与状态一致", "通过"),
        ("IT-04", "超时处理幂等", "重复扫描只标记一次、只扣一次信用并只生成一条异常", "通过"),
        ("IT-05", "并发确认申请", "并发选择仅产生一个 task_order", "通过"),
        ("IT-06", "图片凭证上传", "上传成功并返回可访问 URL", "通过"),
    ], [1.7, 4.2, 8.4, 1.7], 8.5)
    add_heading(doc, "6.3 构建与人工验收记录", 2)
    add_table(doc, ["验收项", "执行命令/操作", "预期结果", "当前结果"], [
        ("后端测试打包", "mvn clean verify", "测试全通过并生成 JAR", "通过"),
        ("Web 生产构建", "npm run build", "dist 资源生成，无编译错误", "通过"),
        ("Android 同步", "npm run android:sync", "Web 产物复制并同步 Capacitor 插件", "通过"),
        ("Android APK", "gradlew.bat assembleDebug", "生成 app-debug.apk", "通过"),
        ("模拟器启动", "Pixel 8 点击 Run", "显示 CampusLink 移动端", "通过"),
        ("模拟器联网", "启动后端后登录 alice/demo123", "10.0.2.2:8080 请求成功", "通过"),
        ("返回导航", "打开任务详情/发布页/订单页", "出现左箭头且返回语义正确", "通过"),
    ], [3.0, 4.5, 5.5, 2.0], 8.5)
    add_heading(doc, "6.4 核心验收用例", 2)
    add_table(doc, ["编号", "操作步骤", "期望结果", "优先级"], [
        ("AT-01", "以 alice 登录，发布奖励 20 分任务", "发布成功；alice 可用积分减少 20、冻结积分增加 20", "P0"),
        ("AT-02", "以 bob 申请任务，alice 接受申请", "仅生成一个订单；其他申请变为 REJECTED", "P0"),
        ("AT-03", "bob 开始并提交凭证，alice 验收", "状态依次变化；20 分结算给 bob；双方收到通知", "P0"),
        ("AT-04", "alice 在招募期取消本人任务", "任务 CANCELLED；冻结积分退回", "P0"),
        ("AT-05", "任务双方发起异常，admin 选择 REFUND/SETTLE/RESUME", "裁决仅执行一次；积分、信用与状态一致", "P0"),
        ("AT-06", "维护技能供需并查看匹配", "返回最多 10 人，显示匹配技能、分数和双向互补", "P1"),
        ("AT-07", "普通用户直接进入 /admin 或调用管理 API", "前端重定向；后端返回 403", "P0"),
        ("AT-08", "在 390px 宽度与 Pixel 8 模拟器浏览核心页面", "无横向溢出，导航、表单和返回操作可用", "P1"),
    ], [1.7, 7.0, 5.8, 1.5], 8.5)

    page_break(doc)
    add_heading(doc, "7 部署、运行与运维", 1)
    add_heading(doc, "7.1 本地快速运行", 2)
    add_body(doc, "快速演示不需要预先安装 MySQL。后端默认使用 H2 内存数据库并自动加载 admin、alice、bob 演示账户；前端 Vite 开发服务器将 /api 和 /uploads 转发到 8080。应分别保持后端和前端终端运行。")
    add_code(doc, "# 终端 1：后端\nSet-Location D:\\CampusLink\\backend\nmvn spring-boot:run\n\n# 终端 2：Web 前端\nSet-Location D:\\CampusLink\\frontend\nnpm install\nnpm run dev")
    add_table(doc, ["角色", "用户名", "密码", "用途"], [
        ("管理员", "admin", "admin123", "治理后台与异常裁决"),
        ("学生/发布者", "alice", "demo123", "发布任务和验收"),
        ("学生/接取者", "bob", "demo123", "申请、履约和技能互助"),
    ], [3.0, 3.0, 3.0, 7.0])
    add_heading(doc, "7.2 Docker Compose 部署", 2)
    add_picture(doc, deployment, 6.7, "图 7-1  当前三服务部署拓扑")
    add_code(doc, "Set-Location D:\\CampusLink\nCopy-Item .env.example .env\n# 编辑 .env：设置 MYSQL_ROOT_PASSWORD、MYSQL_PASSWORD、CAMPUSLINK_TOKEN_SECRET\ndocker compose up --build -d")
    add_body(doc, "默认 Web 地址为 http://localhost，API 为 http://localhost:8080，MySQL 为 localhost:3306。生产环境应设置 CAMPUSLINK_SEED_DEMO_DATA=false，配置高强度令牌密钥，并通过首启环境变量创建管理员。对外部署必须为前后端配置 HTTPS，不应公开数据库端口。")
    add_heading(doc, "7.3 Android 模拟器运行", 2)
    add_code(doc, "# 先保持后端运行，再同步 Android 工程\nSet-Location D:\\CampusLink\\frontend\nnpm install\nnpm run android:sync\nnpm run android:open\n\n# Android Studio：选择 JVM 21、app 配置、Pixel 8，然后点击 Run")
    add_body(doc, "模拟器不能用 localhost 访问开发电脑；Android 官方模拟器把宿主机映射为 10.0.2.2。若显示 Network Error，应依次确认后端日志已出现 Tomcat started on port 8080、frontend/.env.android 配置正确、执行过 android:sync，并确认防火墙或 VPN 未拦截本地连接。VPN 端口 7897 通常不应影响 10.0.2.2，但全局代理或网络接管模式可能需要临时关闭后复测。")
    add_heading(doc, "7.4 主要配置项", 2)
    add_table(doc, ["变量", "默认值", "说明"], [
        ("SPRING_DATASOURCE_URL", "H2 内存库", "部署时改为 MySQL JDBC 地址"),
        ("SPRING_DATASOURCE_USERNAME", "sa", "数据库用户名"),
        ("SPRING_DATASOURCE_PASSWORD", "空", "数据库密码"),
        ("CAMPUSLINK_TOKEN_SECRET", "开发默认值", "生产必须替换为高强度随机密钥"),
        ("CAMPUSLINK_SEED_DEMO_DATA", "true", "是否加载演示账户和样例数据"),
        ("CAMPUSLINK_BOOTSTRAP_ADMIN_USERNAME", "空", "关闭演示数据后首个管理员用户名"),
        ("CAMPUSLINK_BOOTSTRAP_ADMIN_PASSWORD", "空", "首个管理员密码，至少 8 位"),
        ("CAMPUSLINK_UPLOAD_DIR", "./uploads", "图片持久化目录"),
        ("CAMPUSLINK_TIMEOUT_CRON", "0 */10 * * * *", "超时扫描 Cron"),
        ("VITE_API_BASE", "Web 留空/Android 为 10.0.2.2:8080", "前端 API 根地址"),
    ], [5.5, 4.2, 6.3], 8.5)
    add_heading(doc, "7.5 运维与安全要求", 2)
    add_bullets(doc, [
        "禁止提交 .env、Android local.properties、签名密钥、生产令牌密钥及用户上传文件。",
        "定期备份 campuslink_mysql 与 campuslink_uploads；恢复时应保持数据库与上传目录时间点一致。",
        "数据库升级只新增 Flyway 迁移，不修改已在线执行过的 V1～V4 文件。",
        "生产环境关闭演示数据、修改默认管理员密码、配置 HTTPS，并限制 MySQL 端口只对内网开放。",
        "监控后端 5xx、鉴权失败、超时任务数量、异常裁决积压和磁盘上传目录容量。",
        "上传接口当前按图片场景提供，正式运营应增加病毒扫描、对象存储、内容审核和访问签名。",
    ])

    page_break(doc)
    add_heading(doc, "8 用户操作说明", 1)
    add_heading(doc, "8.1 学生用户操作流程", 2)
    add_table(doc, ["场景", "操作路径", "注意事项"], [
        ("注册登录", "注册页填写账户资料→登录页输入用户名/邮箱和密码", "首次注册初始化积分和信用；密码由 BCrypt 摘要保存"),
        ("发布任务", "任务→发布任务→填写标题、分类、说明、地点、时间、截止、奖励和信用要求", "发布时立即冻结奖励积分；余额不足不能发布"),
        ("申请任务", "任务广场→任务详情→填写留言和预计完成时间→申请", "不能申请本人任务；信用不足或重复申请会被拒绝"),
        ("选择接取者", "我的→我发布的任务→详情→查看申请→选择", "任务只允许一名接取者；选择后不可直接取消"),
        ("履约验收", "我的→订单→开始/提交；发布者在订单中验收或驳回", "提交需说明，可上传图片凭证；验收后积分结算"),
        ("技能互助", "技能→维护供给和需求→查看推荐→发起互助", "对方依次接受、开始、完成；双方完成后可评价"),
        ("积分信用", "积分→查看余额、冻结积分、积分流水和信用流水", "流水由业务自动生成，用户不可修改"),
        ("消息与资料", "消息→标记已读；我的→编辑资料", "令牌失效后需重新登录"),
        ("举报异常", "相关业务页→举报/发起异常", "履约异常将暂停任务并等待管理员裁决"),
    ], [3.0, 8.2, 4.8], 8.5)
    add_heading(doc, "8.2 管理员操作流程", 2)
    add_bullets(doc, [
        "使用 ADMIN 角色登录后进入管理台，查看学生数、招募中/进行中/已完成任务和待处理举报统计。",
        "在用户列表按状态分页查询，对违规账户执行启用或禁用；接口同时在后端校验角色。",
        "查看平台任务与普通举报，填写处理结果后提交，处理人和处理时间写入数据库。",
        "异常任务根据证据选择 REFUND、SETTLE 或 RESUME，必要时指定责任用户并调整信用。",
        "维护技能分类，新增或停用分类；停用项不再出现在公共技能列表中。",
    ])
    add_heading(doc, "8.3 常见问题处理", 2)
    add_table(doc, ["现象", "可能原因", "处理方法"], [
        ("Android 显示 Network Error", "后端未启动、API 地址仍为 localhost、未同步资源、VPN/防火墙拦截", "确认 8080 日志；检查 .env.android；重新 android:sync；再排查代理"),
        ("Gradle JVM 不兼容", "Android Studio 默认使用过高 JVM", "选择 JVM 21；当前 Gradle 8.14.3 支持该版本"),
        ("Android 页面仍是旧内容", "Vite 产物未复制到原生工程", "执行 npm run android:sync 后重新 Run"),
        ("登录后 401", "令牌过期或本地存储残留", "退出登录并重新登录；核对系统时间"),
        ("H2 数据重启后消失", "默认使用内存数据库", "需要持久化时使用 Docker Compose/MySQL"),
        ("Docker 启动失败", ".env 缺少必填密码或端口被占用", "复制 .env.example 并填写；调整 FRONTEND_PORT/MYSQL_PORT"),
        ("图片无法访问", "上传目录无权限或卷未挂载", "检查 CAMPUSLINK_UPLOAD_DIR、campuslink_uploads 与 /uploads 代理"),
    ], [4.0, 5.4, 6.6], 8.5)

    page_break(doc)
    add_heading(doc, "9 需求追踪、限制与后续计划", 1)
    add_heading(doc, "9.1 需求—实现—测试追踪", 2)
    add_table(doc, ["需求模块", "主要实现", "验证方式", "状态"], [
        ("账户与个人中心", "Auth/Account 服务、令牌、资料页、角色守卫", "登录注册与鉴权接口、人工页面测试", "已完成"),
        ("微任务管理", "发布、检索、申请、确认与取消", "IT-01、IT-05、AT-01～04", "已完成"),
        ("任务履约", "开始、提交、验收、驳回与状态同步", "IT-01、AT-03", "已完成"),
        ("技能互助与匹配", "供需档案、加权推荐、互助状态机", "IT-02、AT-06", "已完成"),
        ("积分信用评价", "冻结、退款、结算、流水、评价信用联动", "IT-01、IT-03、AT-01/03", "已完成"),
        ("举报与平台治理", "普通举报、异常任务、管理员列表与裁决", "IT-03、IT-04、AT-05/07", "已完成"),
        ("Web 响应式", "CSS 断点、导航与卡片重排", "Vite 构建、390px 人工验收", "已完成"),
        ("Android 终端", "Capacitor 工程、Hash 路由、模拟器 API 地址", "APK 构建、Pixel 8 启动与登录", "已完成"),
        ("持久化部署", "MySQL/Flyway、Docker Compose、上传卷", "Compose 配置与 MySQL 运行", "已完成"),
    ], [3.4, 6.2, 4.8, 1.8], 8.5)
    add_heading(doc, "9.2 已知限制", 2)
    add_bullets(doc, [
        "当前认证未集成校园统一身份认证、短信验证码、邮件验证和密码找回；演示账号仅用于课程环境。",
        "消息为站内通知和轮询式读取，没有 WebSocket 实时会话、推送通知或离线消息通道。",
        "文件存储为本地目录/命名卷，没有对象存储、缩略图、恶意文件扫描与内容审核。",
        "技能推荐为可解释规则算法，尚未引入地理距离、时间冲突优化、协同过滤或机器学习模型。",
        "课程版本未实施限流、分布式缓存、链路追踪、指标告警和多实例会话管理。",
        "自动化测试以核心后端闭环为主，前端组件测试、E2E、Android UI 自动化和性能压测仍需补充。",
    ])
    add_heading(doc, "9.3 后续迭代建议", 2)
    add_table(doc, ["优先级", "迭代项", "目标"], [
        ("P0", "生产安全加固", "HTTPS、强密钥、验证码/限流、密码找回、审计日志、上传安全"),
        ("P0", "测试体系", "补充 Vitest、Playwright、Android Espresso、MySQL Testcontainers 与并发压测"),
        ("P1", "消息能力", "WebSocket/推送通知、会话与未读聚合"),
        ("P1", "文件与证据", "对象存储、签名 URL、图片压缩、内容审核和证据保全"),
        ("P1", "推荐优化", "加入时间地点、偏好、完成率与行为反馈，并持续评估可解释性"),
        ("P2", "校园集成", "统一身份认证、校园地图、教务/社团接口与运营后台"),
        ("P2", "可观测与扩展", "Redis 缓存、指标监控、集中日志、容器编排和灰度发布"),
    ], [2.0, 4.3, 9.7], 8.5)

    page_break(doc)
    add_heading(doc, "附录 A 工程目录与构建产物", 1)
    add_code(doc, "CampusLink/\n├─ backend/\n│  ├─ src/main/java/com/campuslink/  控制器、服务、映射、领域对象与配置\n│  ├─ src/main/resources/db/migration/  V1～V4 数据库迁移\n│  ├─ src/test/  六个核心集成测试\n│  └─ pom.xml\n├─ frontend/\n│  ├─ src/  Vue 页面、路由、状态、API 与样式\n│  ├─ android/  Capacitor Android 原生工程\n│  ├─ capacitor.config.json\n│  └─ package.json\n├─ docker-compose.yml\n├─ README.md\n└─ DEVELOPMENT_PROGRESS.md")
    add_table(doc, ["构建对象", "命令", "产物/入口"], [
        ("后端", "mvn clean verify", "backend/target/campuslink-backend-1.0.0.jar"),
        ("Web", "npm run build", "frontend/dist/"),
        ("Android", "npm run android:sync；gradlew.bat assembleDebug", "frontend/android/app/build/outputs/apk/debug/app-debug.apk"),
        ("接口文档", "启动后端", "http://localhost:8080/swagger-ui.html"),
        ("完整部署", "docker compose up --build -d", "Web :80、API :8080、MySQL :3306"),
    ], [3.0, 7.0, 6.0], 8.5)

    add_heading(doc, "附录 B 文档结论", 1)
    add_body(doc, "CampusLink 当前版本已形成从需求、设计、实现到测试和部署的完整课程设计成果。系统核心价值在于以积分冻结/结算保障任务激励，以信用和评价约束履约，以技能供需匹配促进校园互助，并通过异常任务、超时幂等和管理员裁决补齐治理闭环。Web 与 Android 共用业务代码和 REST API，后端可在 H2 演示模式或 MySQL 持久化模式运行，具备继续扩展为校园实际服务的工程基础。")

    # 统一新增章节的默认样式，并保留原模板中的目录、页眉页脚和浮动图形。
    for style_name in ("Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        style = doc.styles[style_name]
        style.font.name = "微软雅黑"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.color.rgb = RGBColor.from_string(GREEN)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()
